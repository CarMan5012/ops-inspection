from __future__ import annotations

import logging
import re
from collections import defaultdict
from datetime import datetime
from pathlib import Path
from typing import Any
from zoneinfo import ZoneInfo

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.shared import Inches, Pt
from docx.text.paragraph import Paragraph

from app.settings import settings
from app.storage_paths import get_report_dir, resolve_artifact_path
from app.utils import relative_to_data, safe_name

logger = logging.getLogger("app.report")


def normalize_section_name(value: str) -> str:
    return re.sub(r"[\s\u3000:：，,。._\-]+", "", str(value or "")).lower()


def report_template_path() -> Path:
    return Path(__file__).parent.parent / "服务器巡检-yyyy-mm-dd.docx"


def list_template_sections() -> list[str]:
    template_path = report_template_path()
    if not template_path.exists():
        return []

    try:
        doc = Document(template_path)
    except Exception as exc:
        logger.warning(f"读取 Word 模板章节失败: {exc}")
        return []

    sections: list[str] = []
    seen: set[str] = set()
    in_period = False
    for p in doc.paragraphs:
        text = p.text.strip()
        if not text:
            continue
        if text.startswith("巡检时间："):
            in_period = True
            continue
        if not in_period:
            continue
        if text[0].isdigit():
            key = normalize_section_name(text)
            if key and key not in seen:
                sections.append(text)
                seen.add(key)
    return sections


def is_header_paragraph(p) -> bool:
    text = p.text.strip()
    if not text:
        return False
    if p.style.name.startswith("Heading"):
        return True
    if text[0].isdigit():
        return True
    if text.startswith("巡检时间："):
        return True
    return False


def insert_paragraph_after(paragraph, text=None, style=None):
    new_p = OxmlElement('w:p')
    paragraph._p.addnext(new_p)
    new_para = Paragraph(new_p, paragraph._parent)
    if text:
        new_para.text = text
    if style:
        new_para.style = style
    return new_para



def generate_docx(
    job: dict[str, Any],
    run: dict[str, Any],
    results: list[dict[str, Any]],
) -> str:
    output_dir = get_report_dir(run["id"])

    logger.info(f"开始为任务 '{job['name']}' 生成巡检报告 Word 文档 (运行 ID: {run['id']})")

    # 1. Determine period (上午 or 下午)
    local_tz = ZoneInfo(settings.default_timezone)
    now_local = datetime.now(local_tz)
    period = "上午" if now_local.hour < 13 else "下午"
    report_name = safe_name(str(job.get("report_title") or job.get("name") or "巡检报告"))
    file_name = f"{report_name}-{now_local.strftime('%Y-%m-%d')}.docx"
    output_path = output_dir / file_name
    logger.info(f"当前时间为 {now_local.strftime('%Y-%m-%d %H:%M:%S')}，匹配巡检时段: '{period}'")

    # 2. 计算开始时间、结束时间及总耗时
    started_at_str = run.get("started_at") or ""
    finished_at_str = ""
    duration_str = "小于1秒"
    
    if started_at_str:
        try:
            if "T" in started_at_str:
                start_dt = datetime.fromisoformat(started_at_str)
            else:
                start_dt = datetime.strptime(started_at_str, "%Y-%m-%d %H:%M:%S")
            
            end_dt = datetime.now(ZoneInfo(settings.default_timezone))
            finished_at_str = end_dt.strftime("%Y-%m-%d %H:%M:%S")
            delta = end_dt.replace(tzinfo=None) - start_dt.replace(tzinfo=None)
            total_seconds = int(delta.total_seconds())
            if total_seconds < 60:
                duration_str = f"{total_seconds}秒"
            else:
                minutes = total_seconds // 60
                seconds = total_seconds % 60
                duration_str = f"{minutes}分{seconds}秒"
        except Exception as e:
            logger.error(f"计算报告耗时出错: {e}", exc_info=True)
            finished_at_str = datetime.now().strftime("%Y-%m-%d %H:%M:%S")
    else:
        started_at_str = "-"
        finished_at_str = datetime.now().strftime("%Y-%m-%d %H:%M:%S")

    # 3. Find template or previous report
    from app.repository import get_previous_run_today
    prev_run = get_previous_run_today(job["id"], run["id"])
    
    template_path = report_template_path()
    
    source_path = None
    if prev_run:
        prev_report_path = Path(prev_run["report_path"])
        if prev_report_path.exists():
            source_path = prev_report_path
            logger.info(f"查找到当天此任务较早执行记录 (运行 ID: {prev_run['id']})，将基于该报告增量合并: {source_path}")
            
    if not source_path:
        if template_path.exists():
            source_path = template_path
            logger.info(f"未找到当天已存报告，将从模板创建新报告: {source_path}")
        else:
            logger.warning(f"Word 模板文件不存在: {template_path}，将降级从零生成默认样式报告")
            
    if source_path:
        try:
            doc = Document(source_path)
            # 1. 替换标题日期（如果是模板）
            today_cn_str = now_local.strftime("%Y年%m月%d日")
            for p in doc.paragraphs:
                if "yyyy年mm月dd日" in p.text:
                    p.text = p.text.replace("yyyy年mm月dd日", today_cn_str)
                    logger.info(f"已更新模板的主标题日期为: {today_cn_str}")
                        
            # 2. 删除旧版本曾插入的执行明细段落。模板正文只保留巡检项和截图。
            for p in list(doc.paragraphs):
                if "巡检执行时间：" in p.text or "巡检耗时：" in p.text:
                    p._element.getparent().remove(p._element)
                        
            # Edit screenshots
            unmatched_items = []
            for item in results:
                item_name = item.get("item_name") or ""
                section_name = item.get("section") or item_name
                status = item.get("status") or "failed"
                file_path = item.get("file_path") or ""
                error_message = item.get("error_message") or ""

                # Find period heading
                period_idx = -1
                for i, p in enumerate(doc.paragraphs):
                    if "巡检时间：" in p.text and period in p.text:
                        period_idx = i
                        break
                
                if period_idx == -1:
                    logger.warning(f"模板中未找到 '{period}' 巡检时段，将把截图项 '{item_name}' 放入新增部分")
                    unmatched_items.append(item)
                    continue
                    
                # Find item heading under period
                item_idx = -1
                candidates = [section_name, item_name]
                clean_candidates = []
                for candidate in candidates:
                    clean_candidate = normalize_section_name(candidate)
                    if clean_candidate and clean_candidate not in clean_candidates:
                        clean_candidates.append(clean_candidate)

                for i in range(period_idx + 1, len(doc.paragraphs)):
                    p = doc.paragraphs[i]
                    if "巡检时间：" in p.text:
                        break
                    p_text_clean = normalize_section_name(p.text)
                    if not p_text_clean:
                        continue
                    if any(
                        clean_candidate in p_text_clean or p_text_clean in clean_candidate
                        for clean_candidate in clean_candidates
                    ):
                        item_idx = i
                        break
                        
                if item_idx == -1:
                    logger.warning(f"未在 '{period}' 时段下匹配到 Word 章节 '{section_name}'，将把截图项 '{item_name}' 放入新增部分")
                    unmatched_items.append(item)
                    continue
                    
                logger.info(f"正在往时段 '{period}' 的 Word 章节 '{section_name}' 插入截图项 '{item_name}'...")
                
                target_p = doc.paragraphs[item_idx]
                
                # Clear paragraphs between item heading and the next heading
                idx = item_idx + 1
                deleted_count = 0
                while idx < len(doc.paragraphs):
                    p = doc.paragraphs[idx]
                    if is_header_paragraph(p):
                        break
                    p_el = p._element
                    p_el.getparent().remove(p_el)
                    deleted_count += 1
                if deleted_count > 0:
                    logger.info(f"清理了该监控项下原有的 {deleted_count} 个过渡段落或旧图")
                    
                # Insert screenshot
                image_path = resolve_artifact_path(file_path) if file_path else Path("")
                if status == "success" and file_path and image_path.exists():
                    new_p = insert_paragraph_after(target_p)
                    run_el = new_p.add_run()
                    try:
                        run_el.add_picture(str(image_path), width=Inches(6.0))
                        logger.info(f"成功将截图图片写入文档: {image_path}")
                    except Exception as exc:
                        logger.error(f"向 Word 写入图片文件失败: {exc}", exc_info=True)
                        new_p.text = f"图片插入失败：{exc}"
                else:
                    logger.warning(f"截图项 '{item_name}' 截图失败，跳过向 Word 写入任何错误文字")

            success_unmatched = [item for item in unmatched_items if item.get("status") == "success"]
            if success_unmatched:
                doc.add_heading("新增巡检截图", level=1)
                for item in success_unmatched:
                    item_name = item.get("item_name") or ""
                    section_name = item.get("section") or item_name
                    file_path = item.get("file_path") or ""
                    image_path = resolve_artifact_path(file_path) if file_path else Path("")
                    
                    if file_path and image_path.exists():
                        doc.add_heading(section_name, level=2)
                        p = doc.add_paragraph()
                        run_el = p.add_run()
                        try:
                            run_el.add_picture(str(image_path), width=Inches(6.0))
                            logger.info(f"成功将未在模板匹配的截图追加至新增章节: {image_path}")
                        except Exception as exc:
                            logger.error(f"向 Word 写入追加图片失败: {exc}", exc_info=True)
                            p.text = f"图片插入失败：{exc}"

            doc.save(output_path)
            logger.info(f"新报告已成功保存至: {output_path}")
            return str(output_path)
        except Exception as exc:
            logger.error(f"基于模板生成报告发生异常: {exc}，将降级使用从零生成报告方案", exc_info=True)
            
    # Fallback to generating from scratch
    doc = Document()
    _setup_styles(doc)
    _add_cover(doc, job, run, finished_at_str, duration_str)
    _add_summary(doc, job, run, results)
    _add_screenshots(doc, results)
    _add_appendix(doc, results)
    doc.save(output_path)
    logger.info(f"降级报告已保存至: {output_path}")
    return str(output_path)



def _setup_styles(doc: Document) -> None:
    normal = doc.styles["Normal"]
    normal.font.name = "Microsoft YaHei"
    normal.font.size = Pt(10.5)
    for style_name in ("Title", "Heading 1", "Heading 2"):
        style = doc.styles[style_name]
        style.font.name = "Microsoft YaHei"


def _add_cover(
    doc: Document,
    job: dict[str, Any],
    run: dict[str, Any],
    finished_at_str: str,
    duration_str: str,
) -> None:
    title = str(job.get("report_title") or "自动化巡检报告")
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run_title = p.add_run(title)
    run_title.bold = True
    run_title.font.size = Pt(24)

    doc.add_paragraph("")
    table = doc.add_table(rows=4, cols=2)
    table.style = "Table Grid"
    fields = [
        ("巡检任务", job.get("name", "")),
        ("巡检环境", job.get("environment", "")),
        ("时间范围", job.get("time_range_label", "")),
        ("执行状态", run.get("status", "")),
    ]
    for idx, (key, value) in enumerate(fields):
        table.cell(idx, 0).text = str(key)
        table.cell(idx, 1).text = str(value or "")
    doc.add_page_break()


def _add_summary(
    doc: Document,
    job: dict[str, Any],
    run: dict[str, Any],
    results: list[dict[str, Any]],
) -> None:
    success_count = len([item for item in results if item["status"] == "success"])
    failed_count = len([item for item in results if item["status"] != "success"])
    if not results:
        conclusion = "未执行截图"
    else:
        conclusion = "正常" if failed_count == 0 else "需关注"

    doc.add_heading("一、巡检结论", level=1)
    doc.add_paragraph(f"总体结论：{conclusion}")
    doc.add_paragraph(f"任务名称：{job.get('name', '')}")
    doc.add_paragraph(f"巡检环境：{job.get('environment', '')}")
    doc.add_paragraph(f"时间范围：{job.get('time_range_label', '')}")

    table = doc.add_table(rows=4, cols=2)
    table.style = "Table Grid"
    fields = [
        ("截图总数", len(results)),
        ("成功截图数", success_count),
        ("失败截图数", failed_count),
        ("报告路径", relative_to_data(run.get("report_path", "")) if run.get("report_path") else ""),
    ]
    for idx, (key, value) in enumerate(fields):
        table.cell(idx, 0).text = str(key)
        table.cell(idx, 1).text = str(value)


def _add_screenshots(doc: Document, results: list[dict[str, Any]]) -> None:
    grouped: dict[str, list[dict[str, Any]]] = defaultdict(list)
    for item in results:
        grouped[item.get("section") or "巡检截图"].append(item)

    doc.add_heading("二、巡检截图", level=1)
    if not results:
        doc.add_paragraph("本次运行没有截图结果，因此文档中没有截图内容。请先在任务中添加并启用截图项，然后点击任务的“立即运行”生成正式报告。")
        return

    for section, items in grouped.items():
        success_items = [it for it in items if it.get("status") == "success"]
        if not success_items:
            continue
            
        doc.add_heading(section, level=2)
        for item in success_items:
            doc.add_paragraph(str(item.get("item_name") or ""), style=None)
            file_path = item.get("file_path") or ""
            image_path = resolve_artifact_path(file_path) if file_path else Path("")
            if file_path and image_path.exists():
                try:
                    doc.add_picture(str(image_path), width=Inches(6.6))
                except Exception as exc:  # noqa: BLE001 - keep report generation alive.
                    doc.add_paragraph(f"图片插入失败：{exc}")
            doc.add_paragraph("")


def _add_appendix(doc: Document, results: list[dict[str, Any]]) -> None:
    doc.add_heading("三、附录", level=1)
    table = doc.add_table(rows=1, cols=4)
    table.style = "Table Grid"
    headers = ("截图名称", "状态", "文件路径", "错误信息")
    for idx, header in enumerate(headers):
        table.cell(0, idx).text = header
    for item in results:
        row = table.add_row().cells
        row[0].text = str(item.get("item_name") or "")
        row[1].text = str(item.get("status") or "")
        row[2].text = relative_to_data(item.get("file_path") or "")
        row[3].text = str(item.get("error_message") or "")
