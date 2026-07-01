from __future__ import annotations

from pathlib import Path
from typing import Any

from fastapi import APIRouter, BackgroundTasks, Depends, FastAPI, Form, HTTPException, Request
from fastapi.openapi.docs import get_redoc_html, get_swagger_ui_html
from fastapi.responses import FileResponse, HTMLResponse, RedirectResponse
from fastapi.staticfiles import StaticFiles
from fastapi.templating import Jinja2Templates

from app.auth import COOKIE_NAME, check_password, create_token, get_session_ttl_seconds, require_login, verify_token
from app.db import init_db
from app.repository import (
    create_run,
    delete_auth_profile,
    delete_job,
    delete_mail_profile,
    delete_screenshot_item,
    get_auth_profile,
    get_job,
    get_mail_profile,
    get_run,
    get_screenshot_item,
    list_auth_profiles,
    list_jobs,
    list_mail_profiles,
    list_runs,
    list_screenshot_items,
    list_screenshot_results,
    save_auth_profile,
    save_job,
    save_mail_profile,
    save_screenshot_item,
    get_system_setting,
    set_system_setting,
)
from app.runner import run_job, test_capture_item
from app.scheduler import reload_jobs, shutdown_scheduler, start_scheduler
from app.settings import settings
from app.storage_paths import get_run_date, resolve_artifact_path
from app.report import list_template_sections
from app.openapi import configure_openapi
from app.login_crypto import LoginPayloadError, decrypt_login_payload, get_login_public_jwk
import os
import logging
import sys
from logging.handlers import RotatingFileHandler


# 获取并规范化 LOG_LEVEL
log_level_str = os.getenv("LOG_LEVEL", "INFO").upper().strip()
log_level = getattr(logging, log_level_str, logging.INFO)

app_logger = logging.getLogger("app")
app_logger.setLevel(log_level)

if not app_logger.handlers:
    formatter = logging.Formatter(
        "%(asctime)s [%(levelname)s] %(name)s: %(message)s",
        datefmt="%Y-%m-%d %H:%M:%S"
    )
    
    # 控制台标准输出
    stream_handler = logging.StreamHandler(sys.stdout)
    stream_handler.setFormatter(formatter)
    app_logger.addHandler(stream_handler)
    
    # 文件持久化滚动日志
    try:
        settings.ensure_dirs()
        log_file = settings.log_dir / "app.log"
        file_handler = RotatingFileHandler(
            log_file,
            maxBytes=10 * 1024 * 1024,  # 10MB
            backupCount=3,
            encoding="utf-8"
        )
        file_handler.setFormatter(formatter)
        app_logger.addHandler(file_handler)
        app_logger.info(f"日志滚动文件处理器初始化成功. 文件路径: {log_file}")
    except Exception as e:
        # 避免写文件失败导致应用无法启动，做降级处理
        stream_handler.stream.write(f"警告: 初始化文件日志失败: {e}\n")
        
    app_logger.propagate = False
    app_logger.info(f"应用日志系统初始化完成. 日志级别: {log_level_str}")

app = FastAPI(
    title=settings.app_name,
    docs_url=None,
    redoc_url=None,
    openapi_url=None,
    openapi_tags=[],
)
configure_openapi(app, settings.api_prefix, settings.app_name)

api_router = APIRouter(prefix=settings.api_prefix)

templates = Jinja2Templates(directory=["app/templates", "app/static/frontend"])
app.mount("/static", StaticFiles(directory="app/static"), name="static")
if settings.frontend_base_path != "/":
    app.mount(f"{settings.frontend_base_path}/static", StaticFiles(directory="app/static"), name="static_prefix")


def url_for_frontend(path: str) -> str:
    base = settings.frontend_base_path
    if base == "/":
        if not path.startswith("/"):
            return "/" + path
        return path
    
    cleaned_path = path.lstrip("/")
    if not cleaned_path:
        return base + "/"
    if cleaned_path.startswith("#") or cleaned_path.startswith("?"):
        return base + "/" + cleaned_path
    return base + "/" + cleaned_path


def frontend_redirect(path: str, status_code: int = 303) -> RedirectResponse:
    return RedirectResponse(url_for_frontend(path), status_code=status_code)


templates.env.globals["url_for_frontend"] = url_for_frontend

frontend_router = APIRouter(prefix=settings.frontend_base_path if settings.frontend_base_path != "/" else "")


def status_label(value: Any) -> str:
    labels = {
        "running": "运行中",
        "success": "成功",
        "partial_success": "部分成功",
        "failed": "失败",
    }
    return labels.get(str(value or ""), str(value or "-"))


def mail_status_label(value: Any) -> str:
    raw = str(value or "").strip()
    if not raw:
        return "-"
    labels = {
        "skipped": "未发送",
        "sent": "已发送",
        "success": "已发送",
    }
    if raw.startswith("failed:"):
        return "发送失败：" + raw.removeprefix("failed:").strip()
    return labels.get(raw, raw)


def dingtalk_status_label(value: Any) -> str:
    raw = str(value or "").strip()
    if not raw:
        return "-"
    labels = {
        "skipped": "未发送",
        "sent": "已发送",
        "success": "已发送",
    }
    if raw.startswith("failed:"):
        return "发送失败：" + raw.removeprefix("failed:").strip()
    return labels.get(raw, raw)


def short_date(value: Any) -> str:
    if not value:
        return "-"
    return str(value).replace("T", " ").split("+")[0][:19]


templates.env.filters["status_label"] = status_label
templates.env.filters["mail_status_label"] = mail_status_label
templates.env.filters["dingtalk_status_label"] = dingtalk_status_label
templates.env.filters["short_date"] = short_date

from app.utils import (
    build_totp_uri,
    decrypt_secret,
    generate_totp_secret,
    mask_value,
    resolve_auth_credential,
    resolve_mail_credential,
    verify_totp_code,
)

templates.env.globals["mask_value"] = mask_value
templates.env.globals["resolve_auth_credential"] = resolve_auth_credential
templates.env.globals["resolve_mail_credential"] = resolve_mail_credential
templates.env.globals["decrypt_secret"] = decrypt_secret


def job_form(
    name: str = Form(...),
    environment: str = Form("生产环境"),
    enabled: str | None = Form(None),
    cron_expression: str = Form("0 9 * * *"),
    time_range_label: str = Form("最近24小时"),
    report_title: str = Form("自动化巡检报告"),
    mail_profile_id: str | None = Form(None),
    send_mail: str | None = Form(None),
    send_mail_on_complete: str | None = Form(None),
    send_mail_on_error: str | None = Form(None),
    browser_width: int = Form(1920),
    browser_height: int = Form(1080),
    headless: str | None = Form(None),
    schedule_mode: str = Form("simple"),
    frequency: str = Form("daily"),
    morning_enabled: str | None = Form(None),
    morning_time: str = Form("09:05"),
    afternoon_enabled: str | None = Form(None),
    afternoon_time: str = Form("17:05"),
    dingtalk_enabled: str | None = Form(None),
    dingtalk_webhook: str | None = Form(""),
    dingtalk_secret: str | None = Form(""),
    dingtalk_keyword: str | None = Form(""),
) -> dict[str, Any]:
    from app.schedule_utils import parse_simple_inspection_schedule
    
    def parse_bool(v: Any) -> bool:
        return bool(v) if v is not None else False
        
    payload = {
        "frequency": frequency,
        "morning_enabled": parse_bool(morning_enabled),
        "morning_time": morning_time,
        "afternoon_enabled": parse_bool(afternoon_enabled),
        "afternoon_time": afternoon_time
    }
    
    if schedule_mode == "simple":
        try:
            cron_expr, label_expr, schedule_config = parse_simple_inspection_schedule(payload)
        except Exception as e:
            raise ValueError(str(e))
    else:
        cron_expr = cron_expression
        label_expr = "高级 Cron"
        schedule_config = ""
        
    send_comp = parse_bool(send_mail_on_complete)
    if send_mail is not None and send_mail_on_complete is None:
        send_comp = parse_bool(send_mail)
    send_err = parse_bool(send_mail_on_error) if send_mail_on_error is not None else True
    
    return {
        "name": name,
        "environment": environment,
        "enabled": 1 if enabled else 0,
        "cron_expression": cron_expr,
        "time_range_label": time_range_label,
        "report_title": report_title,
        "mail_profile_id": parse_optional_int(mail_profile_id),
        "send_mail": 1 if send_comp else 0,
        "send_mail_on_complete": 1 if send_comp else 0,
        "send_mail_on_error": 1 if send_err else 0,
        "browser_width": browser_width,
        "browser_height": browser_height,
        "headless": 1,
        "schedule_mode": schedule_mode,
        "schedule_label": label_expr,
        "schedule_config": schedule_config,
        "dingtalk_enabled": 1 if dingtalk_enabled else 0,
        "dingtalk_webhook": (dingtalk_webhook or "").strip(),
        "dingtalk_secret": (dingtalk_secret or "").strip(),
        "dingtalk_keyword": (dingtalk_keyword or "").strip(),
    }


def item_form(
    auth_profile_id: str | None = Form(None),
    name: str = Form(...),
    item_type: str = Form("web"),
    url: str = Form(...),
    section: str = Form("巡检截图"),
    capture_mode: str = Form("full_page"),
    css_selector: str = Form(""),
    wait_selector: str = Form(""),
    wait_seconds: float = Form(3),
    timeout_seconds: int = Form(60),
    retry_count: int = Form(2),
    browser_width: str | None = Form(None),
    browser_height: str | None = Form(None),
    sort_order: int = Form(100),
    enabled: str | None = Form(None),
    real_browser_capture: str | None = Form(None),
    watermark_enabled: str | None = Form(None),
    watermark_text: str = Form(""),
    watermark_opacity: int = Form(65),
    watermark_font_size: int = Form(24),
    watermark_gap_x: int = Form(140),
    watermark_gap_y: int = Form(140),
    watermark_angle: int = Form(-45),
) -> dict[str, Any]:
    real_capture_enabled = 1 if real_browser_capture else 0
    normalized_capture_mode = capture_mode
    normalized_css_selector = css_selector
    if real_capture_enabled:
        normalized_capture_mode = "viewport"
        normalized_css_selector = ""

    return {
        "job_id": None,
        "auth_profile_id": parse_optional_int(auth_profile_id),
        "name": name,
        "item_type": item_type,
        "url": url,
        "section": section,
        "capture_mode": normalized_capture_mode,
        "css_selector": normalized_css_selector,
        "wait_selector": wait_selector,
        "wait_seconds": wait_seconds,
        "timeout_seconds": timeout_seconds,
        "retry_count": retry_count,
        "browser_width": parse_optional_int(browser_width),
        "browser_height": parse_optional_int(browser_height),
        "sort_order": sort_order,
        "enabled": 1 if enabled else 0,
        "real_browser_capture": real_capture_enabled,
        "watermark_enabled": 1 if watermark_enabled else 0,
        "watermark_text": watermark_text,
        "watermark_opacity": max(0, min(255, int(watermark_opacity))),
        "watermark_font_size": max(10, min(72, int(watermark_font_size))),
        "watermark_gap_x": max(20, min(800, int(watermark_gap_x))),
        "watermark_gap_y": max(20, min(800, int(watermark_gap_y))),
        "watermark_angle": max(-90, min(90, int(watermark_angle))),
    }


def auth_form(
    name: str = Form(...),
    auth_type: str = Form("form"),
    login_url: str = Form(""),
    username_value: str = Form(""),
    password: str = Form(""),
    username_selector: str = Form('input[name="user"], input[name="username"], input[type="email"], input[placeholder*="username"], input[placeholder*="email"], input[placeholder*="user"]'),
    password_selector: str = Form('input[name="password"], input[type="password"], input[placeholder*="password"]'),
    submit_selector: str = Form('button[type="submit"], button:has-text("Log in"), button:has-text("Login"), button:has-text("登录"), input[type="submit"]'),
    success_selector: str = Form(""),
    storage_state_path: str = Form(""),
) -> dict[str, Any]:
    return {
        "name": name,
        "auth_type": auth_type,
        "login_url": login_url,
        "username": username_value,
        "password": password,
        "username_value": username_value,
        "username_source": "plain",
        "password_source": "plain",
        "username_env": "",
        "password_env": "",
        "username_selector": username_selector,
        "password_selector": password_selector,
        "submit_selector": submit_selector,
        "success_selector": success_selector,
        "storage_state_path": storage_state_path,
    }


def mail_form(
    name: str = Form(...),
    smtp_host: str = Form(""),
    smtp_port: int = Form(465),
    use_ssl: str | None = Form(None),
    use_starttls: str | None = Form(None),
    username: str = Form(""),
    password: str = Form(""),
    sender: str = Form(""),
    recipients: str = Form(""),
    cc: str = Form(""),
    subject_template: str = Form("自动化巡检报告 - {date}"),
    body_template: str = Form("巡检报告已生成，请查看附件。"),
) -> dict[str, Any]:
    return {
        "name": name,
        "smtp_host": smtp_host,
        "smtp_port": smtp_port,
        "use_ssl": 1 if use_ssl else 0,
        "use_starttls": 1 if use_starttls else 0,
        "username": username,
        "password": password,
        "password_source": "plain",
        "password_env": "",
        "sender": sender,
        "recipients": recipients,
        "cc": cc,
        "subject_template": subject_template,
        "body_template": body_template,
    }


@app.on_event("startup")
def on_startup() -> None:
    init_db()
    start_scheduler()


@app.on_event("shutdown")
def on_shutdown() -> None:
    shutdown_scheduler()


@frontend_router.get("/login", response_class=HTMLResponse)
def login_page(request: Request) -> HTMLResponse:
    security = security_settings()
    return templates.TemplateResponse(
        "login.html",
        {
            "request": request,
            "error": request.query_params.get("error", ""),
            "mfa_enabled": security["mfa_enabled"],
            "login_public_key": get_login_public_jwk(),
        },
    )


@frontend_router.post("/login")
def login(
    login_payload: str = Form(""),
) -> RedirectResponse:
    if not login_payload:
        return frontend_redirect("/login?error=1")
    try:
        payload = decrypt_login_payload(login_payload)
    except LoginPayloadError:
        return frontend_redirect("/login?error=1")
    username = payload["username"]
    password = payload["password"]
    mfa_code = payload["mfa_code"]

    if not check_password(username, password):
        return frontend_redirect("/login?error=1")
    security = security_settings()
    if security["mfa_enabled"]:
        secret = current_mfa_secret()
        if not secret or not verify_totp_code(secret, mfa_code):
            return frontend_redirect("/login?error=mfa")
    response = frontend_redirect("/")
    response.set_cookie(
        COOKIE_NAME,
        create_token(username),
        httponly=True,
        samesite="lax",
        max_age=get_session_ttl_seconds(),
    )
    return response


@frontend_router.get("/logout")
def logout() -> RedirectResponse:
    response = frontend_redirect("/login")
    response.delete_cookie(COOKIE_NAME)
    return response


@frontend_router.get("/app-config.json")
def get_app_config() -> dict[str, str]:
    return {
        "apiBase": settings.api_prefix,
        "frontendBase": settings.frontend_base_path,
        "artifactBase": "/artifact"
    }


SWAGGER_SETTING_KEY = "swagger_enabled"
MFA_ENABLED_KEY = "mfa_enabled"
MFA_SECRET_KEY = "mfa_totp_secret"
SESSION_TTL_MINUTES_KEY = "session_ttl_minutes"


def security_settings() -> dict[str, Any]:
    return {
        "mfa_enabled": get_system_setting(MFA_ENABLED_KEY, "0").strip().lower() in {"1", "true", "yes", "on", "enabled"},
        "session_ttl_minutes": max(5, min(int(get_system_setting(SESSION_TTL_MINUTES_KEY, "30") or "30"), 7 * 24 * 60)),
        "mfa_configured": bool(get_system_setting(MFA_SECRET_KEY, "")),
    }


def current_mfa_secret() -> str:
    return decrypt_secret(get_system_setting(MFA_SECRET_KEY, ""))


def swagger_docs_enabled() -> bool:
    return get_system_setting(SWAGGER_SETTING_KEY, "0").strip().lower() in {"1", "true", "yes", "on", "enabled"}


def require_swagger_access(request: Request) -> None:
    require_login(request)
    if not swagger_docs_enabled():
        raise HTTPException(status_code=404, detail="接口文档未开启")


@frontend_router.get("/openapi.json", include_in_schema=False)
def openapi_json(_: None = Depends(require_swagger_access)) -> dict[str, Any]:
    return app.openapi()


@frontend_router.get("/docs", include_in_schema=False)
def swagger_docs(_: None = Depends(require_swagger_access)) -> HTMLResponse:
    return get_swagger_ui_html(
        openapi_url=url_for_frontend("/openapi.json"),
        title=f"{settings.app_name} API 文档",
    )


@frontend_router.get("/redoc", include_in_schema=False)
def redoc_docs(_: None = Depends(require_swagger_access)) -> HTMLResponse:
    return get_redoc_html(
        openapi_url=url_for_frontend("/openapi.json"),
        title=f"{settings.app_name} API 文档",
    )


@frontend_router.get("/swagger", include_in_schema=False)
def swagger_page(_: None = Depends(require_swagger_access)) -> RedirectResponse:
    return frontend_redirect("/docs")


def require_api_login(request: Request) -> None:
    if verify_token(request.cookies.get(COOKIE_NAME)):
        return
    raise HTTPException(status_code=401, detail="Unauthorized")


api_status_label = status_label
api_mail_status_label = mail_status_label


def bool_value(value: Any, default: bool = False) -> bool:
    if value is None:
        return bool(default)
    if isinstance(value, bool):
        return value
    if isinstance(value, (int, float)):
        return value != 0
    return str(value).strip().lower() in {"1", "true", "yes", "on", "enabled"}


def text_value(payload: dict[str, Any], key: str, default: str = "") -> str:
    value = payload.get(key, default)
    if value is None:
        return default
    return str(value).strip()


def int_value(payload: dict[str, Any], key: str, default: int) -> int:
    value = payload.get(key, default)
    if value in (None, ""):
        return default
    try:
        return int(value)
    except (ValueError, TypeError) as e:
        raise HTTPException(status_code=400, detail="参数格式错误") from e


def float_value(payload: dict[str, Any], key: str, default: float) -> float:
    value = payload.get(key, default)
    if value in (None, ""):
        return default
    try:
        return float(value)
    except (ValueError, TypeError) as e:
        raise HTTPException(status_code=400, detail="参数格式错误") from e


def clamp_int(value: Any, default: int, min_value: int, max_value: int) -> int:
    if value in (None, ""):
        return default
    try:
        parsed = int(value)
    except (ValueError, TypeError) as e:
        raise HTTPException(status_code=400, detail="参数格式错误") from e
    return max(min_value, min(max_value, parsed))


def public_auth_profile(profile: dict[str, Any]) -> dict[str, Any]:
    data = dict(profile)
    for key in ("password", "password_secret"):
        data.pop(key, None)
    data["has_password"] = bool(profile.get("password") or profile.get("password_secret"))
    data["has_username"] = bool(profile.get("username") or profile.get("username_value"))
    return data


def public_mail_profile(profile: dict[str, Any]) -> dict[str, Any]:
    data = dict(profile)
    for key in ("password", "password_secret"):
        data.pop(key, None)
    data["has_password"] = bool(profile.get("password") or profile.get("password_secret"))
    return data


def public_job(job: dict[str, Any]) -> dict[str, Any]:
    data = dict(job)
    data["has_dingtalk_webhook"] = bool(job.get("dingtalk_webhook"))
    data["has_dingtalk_secret"] = bool(job.get("dingtalk_secret"))
    data["dingtalk_webhook"] = ""
    data["dingtalk_secret"] = ""
    return data


def public_jobs(jobs: list[dict[str, Any]]) -> list[dict[str, Any]]:
    return [public_job(job) for job in jobs]


@api_router.get("/swagger-settings", dependencies=[Depends(require_api_login)])
def api_get_swagger_settings() -> dict[str, Any]:
    return {"enabled": swagger_docs_enabled()}


@api_router.put("/swagger-settings", dependencies=[Depends(require_api_login)])
def api_update_swagger_settings(payload: dict[str, Any]) -> dict[str, Any]:
    enabled = bool_value(payload.get("enabled"), False)
    set_system_setting(SWAGGER_SETTING_KEY, "1" if enabled else "0")
    return {"ok": True, "enabled": bool(enabled)}


@api_router.get("/security-settings", dependencies=[Depends(require_api_login)])
def api_get_security_settings() -> dict[str, Any]:
    data = security_settings()
    data["mfa_secret"] = ""
    data["mfa_otpauth_url"] = ""
    data["qr_uri"] = ""
    return data


@api_router.put("/security-settings", dependencies=[Depends(require_api_login)])
def api_update_security_settings(payload: dict[str, Any]) -> dict[str, Any]:
    session_minutes = max(5, min(int_value(payload, "session_ttl_minutes", 30), 7 * 24 * 60))
    was_mfa_enabled = security_settings()["mfa_enabled"]
    mfa_enabled = bool_value(payload.get("mfa_enabled"), False)
    verify_code = text_value(payload, "mfa_code", "")
    if verify_code:
        mfa_enabled = True
    current_secret = current_mfa_secret()

    if mfa_enabled and not current_secret:
        raise HTTPException(status_code=400, detail="请先生成并绑定 MFA 认证")

    if mfa_enabled and (verify_code or not was_mfa_enabled):
        if not verify_totp_code(current_secret, verify_code):
            raise HTTPException(status_code=400, detail="MFA 验证码不正确")

    set_system_setting(SESSION_TTL_MINUTES_KEY, str(session_minutes))
    set_system_setting(MFA_ENABLED_KEY, "1" if mfa_enabled else "0")

    return {"ok": True, **security_settings()}


@api_router.post("/security-settings/mfa-secret", dependencies=[Depends(require_api_login)])
def api_generate_mfa_secret() -> dict[str, Any]:
    from app.utils import encrypt_secret

    secret = generate_totp_secret()
    set_system_setting(MFA_SECRET_KEY, encrypt_secret(secret))
    set_system_setting(MFA_ENABLED_KEY, "0")
    otpauth_url = build_totp_uri(secret, settings.admin_username, settings.app_name)
    return {
        "ok": True,
        "mfa_enabled": False,
        "mfa_configured": True,
        "mfa_secret": secret,
        "mfa_otpauth_url": otpauth_url,
        "secret": secret,
        "qr_uri": otpauth_url,
    }


def api_job_payload(payload: dict[str, Any]) -> dict[str, Any]:
    from app.schedule_utils import parse_simple_inspection_schedule

    schedule_mode = text_value(payload, "schedule_mode", "simple")
    cron_expr = text_value(payload, "cron_expression", "0 9 * * *")
    schedule_label = text_value(payload, "schedule_label", "高级 Cron")
    schedule_config = text_value(payload, "schedule_config", "")
    
    if schedule_mode == "simple":
        try:
            cron_expr, schedule_label, schedule_config = parse_simple_inspection_schedule(payload)
        except Exception as e:
            raise HTTPException(status_code=400, detail=str(e))

    send_mail_on_complete = bool_value(payload.get("send_mail_on_complete"), True)
    send_mail_on_error = bool_value(payload.get("send_mail_on_error"), True)
    send_mail = bool_value(payload.get("send_mail"), False)
    if "send_mail" in payload and "send_mail_on_complete" not in payload:
        send_mail_on_complete = send_mail

    return {
        "name": text_value(payload, "name", "未命名巡检任务"),
        "environment": text_value(payload, "environment", "生产环境"),
        "enabled": bool_value(payload.get("enabled"), True),
        "cron_expression": cron_expr,
        "time_range_label": text_value(payload, "time_range_label", "最近 24 小时"),
        "report_title": text_value(payload, "report_title", "自动化巡检报告"),
        "mail_profile_id": parse_optional_int(str(payload.get("mail_profile_id") or "")),
        "send_mail": 1 if send_mail_on_complete else 0,
        "send_mail_on_complete": 1 if send_mail_on_complete else 0,
        "send_mail_on_error": 1 if send_mail_on_error else 0,
        "browser_width": int_value(payload, "browser_width", 1920),
        "browser_height": int_value(payload, "browser_height", 1080),
        "headless": True,
        "schedule_mode": schedule_mode,
        "schedule_label": schedule_label,
        "schedule_config": schedule_config,
        "dingtalk_enabled": bool_value(payload.get("dingtalk_enabled"), False),
        "dingtalk_webhook": text_value(payload, "dingtalk_webhook", ""),
        "dingtalk_secret": text_value(payload, "dingtalk_secret", ""),
        "dingtalk_keyword": text_value(payload, "dingtalk_keyword", ""),
    }


def api_item_payload(payload: dict[str, Any], job_id: int) -> dict[str, Any]:
    real_browser_capture = bool_value(payload.get("real_browser_capture"), True)
    capture_mode = text_value(payload, "capture_mode", "viewport")
    css_selector = text_value(payload, "css_selector", "")
    if real_browser_capture:
        capture_mode = "viewport"
        css_selector = ""

    return {
        "job_id": job_id,
        "auth_profile_id": parse_optional_int(str(payload.get("auth_profile_id") or "")),
        "name": text_value(payload, "name", "未命名截图项"),
        "item_type": text_value(payload, "item_type", "grafana"),
        "url": text_value(payload, "url", ""),
        "section": text_value(payload, "section", "服务器资源"),
        "capture_mode": capture_mode,
        "css_selector": css_selector,
        "wait_selector": text_value(payload, "wait_selector", ""),
        "wait_seconds": float_value(payload, "wait_seconds", 3.0),
        "timeout_seconds": int_value(payload, "timeout_seconds", 60),
        "retry_count": int_value(payload, "retry_count", 2),
        "browser_width": parse_optional_int(str(payload.get("browser_width") or "")),
        "browser_height": parse_optional_int(str(payload.get("browser_height") or "")),
        "sort_order": int_value(payload, "sort_order", 100),
        "enabled": bool_value(payload.get("enabled"), True),
        "real_browser_capture": 1 if real_browser_capture else 0,
        "watermark_enabled": bool_value(payload.get("watermark_enabled"), False),
        "watermark_text": text_value(
            payload,
            "watermark_text",
            "{time}\n省客户服务中心\ndwangchengyi7(王诚毅)",
        ),
        "watermark_opacity": clamp_int(payload.get("watermark_opacity"), 65, 0, 255),
        "watermark_font_size": clamp_int(payload.get("watermark_font_size"), 24, 10, 72),
        "watermark_gap_x": clamp_int(payload.get("watermark_gap_x"), 140, 20, 800),
        "watermark_gap_y": clamp_int(payload.get("watermark_gap_y"), 140, 20, 800),
        "watermark_angle": clamp_int(payload.get("watermark_angle"), -45, -90, 90),
    }


def api_auth_payload(payload: dict[str, Any]) -> dict[str, Any]:
    username = text_value(payload, "username_value", text_value(payload, "username", ""))
    return {
        "name": text_value(payload, "name", "未命名认证"),
        "auth_type": text_value(payload, "auth_type", "form"),
        "login_url": text_value(payload, "login_url", ""),
        "username": username,
        "password": text_value(payload, "password", ""),
        "username_source": text_value(payload, "username_source", "plain"),
        "password_source": text_value(payload, "password_source", "plain"),
        "username_value": username,
        "password_secret": "",
        "username_env": text_value(payload, "username_env", ""),
        "password_env": text_value(payload, "password_env", ""),
        "username_selector": text_value(payload, "username_selector", 'input[name="user"], input[name="username"], input[type="email"], input[placeholder*="username"], input[placeholder*="email"], input[placeholder*="user"]'),
        "password_selector": text_value(payload, "password_selector", 'input[name="password"], input[type="password"], input[placeholder*="password"]'),
        "submit_selector": text_value(payload, "submit_selector", 'button[type="submit"], button:has-text("Log in"), button:has-text("Login"), button:has-text("登录"), input[type="submit"]'),
        "success_selector": text_value(payload, "success_selector", ""),
        "storage_state_path": text_value(payload, "storage_state_path", ""),
    }


def api_mail_payload(payload: dict[str, Any]) -> dict[str, Any]:
    return {
        "name": text_value(payload, "name", "未命名邮箱"),
        "smtp_host": text_value(payload, "smtp_host", ""),
        "smtp_port": int_value(payload, "smtp_port", 465),
        "use_ssl": bool_value(payload.get("use_ssl"), True),
        "use_starttls": bool_value(payload.get("use_starttls"), False),
        "username": text_value(payload, "username", ""),
        "password": text_value(payload, "password", ""),
        "password_source": text_value(payload, "password_source", "plain"),
        "password_secret": "",
        "password_env": text_value(payload, "password_env", ""),
        "sender": text_value(payload, "sender", ""),
        "recipients": text_value(payload, "recipients", ""),
        "cc": text_value(payload, "cc", ""),
        "subject_template": text_value(payload, "subject_template", "自动化巡检报告 - {date}"),
        "body_template": text_value(payload, "body_template", "巡检报告已生成，请查看附件。"),
    }


def clear_auth_state_cache(auth_profile: dict[str, Any]) -> None:
    from app.screenshot import safe_delete_auth_state
    safe_delete_auth_state(auth_profile)


def dashboard_metrics(jobs: list[dict[str, Any]], runs: list[dict[str, Any]]) -> dict[str, int]:
    return {
        "jobs_total": len(jobs),
        "jobs_enabled": sum(1 for job in jobs if int(job.get("enabled") or 0)),
        "runs_success": sum(1 for run in runs if run.get("status") == "success"),
        "runs_failed": sum(1 for run in runs if run.get("status") == "failed"),
        "runs_running": sum(1 for run in runs if run.get("status") == "running"),
    }


def artifact_filename(file_path: str) -> str:
    return str(file_path or "").replace("\\", "/").rsplit("/", 1)[-1]


def get_artifact_url(kind: str, run_id: int | str, filename: str) -> str:
    return url_for_frontend(f"/artifact/{kind}/{run_id}/{filename}")


def enrich_run(run: dict[str, Any]) -> dict[str, Any]:
    data = dict(run)
    data["status_label"] = api_status_label(data.get("status"))
    data["mail_status_label"] = api_mail_status_label(data.get("mail_status"))
    data["dingtalk_status_label"] = dingtalk_status_label(data.get("dingtalk_status"))
    if data.get("report_path"):
        filename = artifact_filename(data["report_path"])
        local_path = resolve_artifact_path(data["report_path"])
        if local_path.exists():
            data["report_url"] = get_artifact_url("reports", data["id"], filename)
    return data


@api_router.get("/session", dependencies=[Depends(require_api_login)])
def api_session() -> dict[str, Any]:
    return {"authenticated": True, "app_name": "自动化巡检控制台"}


@api_router.get("/dashboard", dependencies=[Depends(require_api_login)])
def api_dashboard() -> dict[str, Any]:
    jobs = list_jobs()
    runs = [enrich_run(run) for run in list_runs(12)]
    return {
        "app_name": "自动化巡检控制台",
        "jobs": public_jobs(jobs),
        "runs": runs,
        "metrics": dashboard_metrics(jobs, runs),
        "auth_profiles": [public_auth_profile(profile) for profile in list_auth_profiles()],
        "mail_profiles": [public_mail_profile(profile) for profile in list_mail_profiles()],
    }


@api_router.get("/jobs", dependencies=[Depends(require_api_login)])
def api_list_jobs() -> dict[str, Any]:
    jobs = list_jobs()
    return {"jobs": public_jobs(jobs), "metrics": dashboard_metrics(jobs, list_runs(20))}


@api_router.post("/jobs", dependencies=[Depends(require_api_login)])
def api_create_job(payload: dict[str, Any]) -> dict[str, Any]:
    from app.utils import validate_cron_expression
    import sqlite3

    data = api_job_payload(payload)
    validate_cron_expression(data["cron_expression"])
    try:
        job_id = save_job(data)
    except sqlite3.IntegrityError as e:
        if "UNIQUE constraint failed" in str(e) and "report_jobs.name" in str(e):
            raise HTTPException(status_code=400, detail="任务名称已存在")
        raise HTTPException(status_code=400, detail=f"数据库约束校验失败: {e}")
    reload_jobs()
    return {"ok": True, "job": public_job(require_value(get_job(job_id), "任务不存在"))}


@api_router.get("/jobs/{job_id}", dependencies=[Depends(require_api_login)])
def api_get_job(job_id: int) -> dict[str, Any]:
    from app.utils import get_next_run_times

    job = require_value(get_job(job_id), "任务不存在")
    items = list_screenshot_items(job_id)
    next_runs = []
    if int(job.get("enabled") or 0):
        next_runs = get_next_run_times(job.get("cron_expression") or "", limit=3)
    latest = next((run for run in list_runs(50) if run.get("job_id") == job_id), None)
    return {
        "job": public_job(job),
        "items": items,
        "next_runs": next_runs,
        "last_status": api_status_label(latest.get("status") if latest else ""),
        "mail_profiles": [public_mail_profile(profile) for profile in list_mail_profiles()],
        "auth_profiles": [public_auth_profile(profile) for profile in list_auth_profiles()],
        "template_sections": list_template_sections(),
    }


@api_router.put("/jobs/{job_id}", dependencies=[Depends(require_api_login)])
def api_update_job(job_id: int, payload: dict[str, Any]) -> dict[str, Any]:
    from app.utils import validate_cron_expression
    import sqlite3

    old = require_value(get_job(job_id), "任务不存在")
    data = api_job_payload(payload)
    if not data.get("dingtalk_webhook"):
        data["dingtalk_webhook"] = old.get("dingtalk_webhook") or ""
    if not data.get("dingtalk_secret"):
        data["dingtalk_secret"] = old.get("dingtalk_secret") or ""
    validate_cron_expression(data["cron_expression"])
    try:
        save_job(data, job_id)
    except sqlite3.IntegrityError as e:
        if "UNIQUE constraint failed" in str(e) and "report_jobs.name" in str(e):
            raise HTTPException(status_code=400, detail="任务名称已存在")
        raise HTTPException(status_code=400, detail=f"数据库约束校验失败: {e}")
    reload_jobs()
    return {"ok": True, "job": public_job(require_value(get_job(job_id), "任务不存在"))}


@api_router.delete("/jobs/{job_id}", dependencies=[Depends(require_api_login)])
def api_delete_job(job_id: int) -> dict[str, Any]:
    require_value(get_job(job_id), "任务不存在")
    delete_job(job_id)
    reload_jobs()
    return {"ok": True}


@api_router.post("/jobs/{job_id}/run", dependencies=[Depends(require_api_login)])
def api_run_job(job_id: int, background_tasks: BackgroundTasks) -> dict[str, Any]:
    job = require_value(get_job(job_id), "任务不存在")
    run_id = create_run(job_id, str(job["name"]))
    background_tasks.add_task(run_job, job_id, run_id)
    return {"ok": True, "run_id": run_id}


@api_router.post("/jobs/{job_id}/items", dependencies=[Depends(require_api_login)])
def api_create_item(job_id: int, payload: dict[str, Any]) -> dict[str, Any]:
    require_value(get_job(job_id), "任务不存在")
    data = api_item_payload(payload, job_id)
    item_id = save_screenshot_item(data)
    return {"ok": True, "item": get_screenshot_item(item_id)}


@api_router.put("/items/{item_id}", dependencies=[Depends(require_api_login)])
def api_update_item(item_id: int, payload: dict[str, Any]) -> dict[str, Any]:
    current = require_value(get_screenshot_item(item_id), "截图项不存在")
    data = api_item_payload(payload, int(current["job_id"]))
    save_screenshot_item(data, item_id)
    return {"ok": True, "item": get_screenshot_item(item_id)}


@api_router.delete("/items/{item_id}", dependencies=[Depends(require_api_login)])
def api_delete_item(item_id: int) -> dict[str, Any]:
    require_value(get_screenshot_item(item_id), "截图项不存在")
    delete_screenshot_item(item_id)
    return {"ok": True}


@api_router.post("/items/{item_id}/test", dependencies=[Depends(require_api_login)])
def api_test_item(item_id: int, background_tasks: BackgroundTasks) -> dict[str, Any]:
    require_value(get_screenshot_item(item_id), "截图项不存在")
    background_tasks.add_task(test_capture_item, item_id)
    return {"ok": True}


@api_router.get("/runs", dependencies=[Depends(require_api_login)])
def api_list_runs() -> dict[str, Any]:
    return {"runs": [enrich_run(run) for run in list_runs(settings.max_recent_runs)]}


@api_router.get("/runs/{run_id}", dependencies=[Depends(require_api_login)])
def api_get_run(run_id: int) -> dict[str, Any]:
    run = enrich_run(require_value(get_run(run_id), "运行记录不存在"))
    results = [with_file_url(dict(item)) for item in list_screenshot_results(run_id)]
    return {"run": run, "results": results}


@api_router.get("/auth-profiles", dependencies=[Depends(require_api_login)])
def api_list_auth_profiles() -> dict[str, Any]:
    return {"profiles": [public_auth_profile(profile) for profile in list_auth_profiles()]}


@api_router.post("/auth-profiles", dependencies=[Depends(require_api_login)])
def api_create_auth_profile(payload: dict[str, Any]) -> dict[str, Any]:
    data = api_auth_payload(payload)
    auth_id = save_auth_profile(data)
    return {"ok": True, "profile": public_auth_profile(require_value(get_auth_profile(auth_id), "认证配置不存在"))}


@api_router.put("/auth-profiles/{auth_id}", dependencies=[Depends(require_api_login)])
def api_update_auth_profile(auth_id: int, payload: dict[str, Any]) -> dict[str, Any]:
    old = require_value(get_auth_profile(auth_id), "认证配置不存在")
    data = api_auth_payload(payload)
    changed = any(
        str(data.get(field) or "") != str(old.get(field) or "")
        for field in (
            "auth_type", "login_url", "username_value", "username", 
            "username_selector", "password_selector", "submit_selector", 
            "success_selector", "username_source", "password_source", 
            "username_env", "password_env", "storage_state_path"
        )
    )
    if data.get("password"):
        changed = True
    else:
        data["password_secret"] = old.get("password_secret")
    if not data.get("username_value"):
        data["username"] = old.get("username")
        data["username_value"] = old.get("username_value")
    save_auth_profile(data, auth_id)
    if changed:
        clear_auth_state_cache(old)
    return {"ok": True, "profile": public_auth_profile(require_value(get_auth_profile(auth_id), "认证配置不存在"))}


@api_router.delete("/auth-profiles/{auth_id}", dependencies=[Depends(require_api_login)])
def api_delete_auth_profile(auth_id: int) -> dict[str, Any]:
    old = require_value(get_auth_profile(auth_id), "认证配置不存在")
    clear_auth_state_cache(old)
    delete_auth_profile(auth_id)
    return {"ok": True}


@api_router.post("/auth-profiles/{auth_id}/test", dependencies=[Depends(require_api_login)])
def api_test_auth_profile(auth_id: int) -> dict[str, Any]:
    from app.screenshot import test_auth_profile_login

    require_value(get_auth_profile(auth_id), "认证配置不存在")
    result = test_auth_profile_login(auth_id)
    return result


@api_router.get("/mail-profiles", dependencies=[Depends(require_api_login)])
def api_list_mail_profiles() -> dict[str, Any]:
    return {"profiles": [public_mail_profile(profile) for profile in list_mail_profiles()]}


@api_router.post("/mail-profiles", dependencies=[Depends(require_api_login)])
def api_create_mail_profile(payload: dict[str, Any]) -> dict[str, Any]:
    data = api_mail_payload(payload)
    mail_id = save_mail_profile(data)
    return {"ok": True, "profile": public_mail_profile(require_value(get_mail_profile(mail_id), "邮件配置不存在"))}


@api_router.put("/mail-profiles/{mail_id}", dependencies=[Depends(require_api_login)])
def api_update_mail_profile(mail_id: int, payload: dict[str, Any]) -> dict[str, Any]:
    old = require_value(get_mail_profile(mail_id), "邮件配置不存在")
    data = api_mail_payload(payload)
    if not data.get("password"):
        data["password_secret"] = old.get("password_secret")
    save_mail_profile(data, mail_id)
    return {"ok": True, "profile": public_mail_profile(require_value(get_mail_profile(mail_id), "邮件配置不存在"))}


@api_router.delete("/mail-profiles/{mail_id}", dependencies=[Depends(require_api_login)])
def api_delete_mail_profile(mail_id: int) -> dict[str, Any]:
    require_value(get_mail_profile(mail_id), "邮件配置不存在")
    delete_mail_profile(mail_id)
    return {"ok": True}


@api_router.get("/periodic-reports", dependencies=[Depends(require_api_login)])
def api_periodic_reports() -> dict[str, Any]:
    from app.repository import list_periodic_report_runs, list_periodic_report_settings

    runs = []
    for item in list_periodic_report_runs(50):
        run = dict(item)
        zip_path = run.get("zip_path")
        if zip_path and os.path.exists(zip_path):
            size_bytes = os.path.getsize(zip_path)
            run["size_str"] = f"{size_bytes / 1024:.2f} KB" if size_bytes < 1024 * 1024 else f"{size_bytes / (1024 * 1024):.2f} MB"
            run["filename"] = Path(zip_path).name
            run["download_url"] = url_for_frontend(f"/periodic-reports/download/{run['id']}")
        else:
            run["size_str"] = "-"
            run["filename"] = ""
        runs.append(run)
    return {"settings": list_periodic_report_settings(), "runs": runs}


@api_router.post("/periodic-reports/{report_type}/run", dependencies=[Depends(require_api_login)])
def api_run_periodic_report(report_type: str, background_tasks: BackgroundTasks) -> dict[str, Any]:
    from app.periodic_reporter import trigger_periodic_report

    background_tasks.add_task(trigger_periodic_report, report_type, True)
    return {"ok": True}


@frontend_router.get("/", response_class=HTMLResponse, dependencies=[Depends(require_login)])
def index(request: Request) -> HTMLResponse:
    return templates.TemplateResponse(
        "index.html",
        {
            "request": request,
            "url_for_frontend": url_for_frontend,
            "api_prefix": settings.api_prefix,
            "frontend_base_path": settings.frontend_base_path,
        },
    )


@frontend_router.get("/jobs/new", response_class=HTMLResponse, dependencies=[Depends(require_login)])
def new_job(request: Request) -> HTMLResponse:
    return templates.TemplateResponse(
        "job_form.html",
        {
            "request": request,
            "job": {},
            "mail_profiles": list_mail_profiles(),
            "action": url_for_frontend("/jobs/new"),
            "error": request.query_params.get("error", "")
        },
    )


@frontend_router.post("/jobs/new", dependencies=[Depends(require_login)])
def create_job(form: dict[str, Any] = Depends(job_form)) -> RedirectResponse:
    from app.utils import validate_cron_expression
    from urllib.parse import quote
    import sqlite3
    try:
        validate_cron_expression(form["cron_expression"])
    except ValueError as e:
        return frontend_redirect(f"/jobs/new?error={quote(str(e))}")

    try:
        job_id = save_job(form)
    except sqlite3.IntegrityError as e:
        if "UNIQUE constraint failed" in str(e) and "report_jobs.name" in str(e):
            return frontend_redirect(f"/jobs/new?error={quote('任务名称已存在')}")
        return frontend_redirect(f"/jobs/new?error={quote(f'数据库错误: {e}')}")
    reload_jobs()
    return frontend_redirect(f"/jobs/{job_id}")


@frontend_router.get("/jobs/{job_id}", response_class=HTMLResponse, dependencies=[Depends(require_login)])
def job_detail(request: Request, job_id: int) -> HTMLResponse:
    from app.scheduler import scheduler
    db_job = get_job(job_id)
    job = dict(db_job) if db_job else None
    if not job:
        raise HTTPException(status_code=404, detail="任务不存在")
        
    next_runs = []
    if int(job.get("enabled") or 0):
        from app.utils import get_next_run_times
        next_runs = get_next_run_times(job.get("cron_expression") or "", limit=3)
    
    items = list_screenshot_items(job_id)
    
    runs_list = list_runs(50)
    last_status = "暂无运行"
    for run in runs_list:
        if run.get("job_id") == job_id:
            from app.main import status_label
            last_status = status_label(run.get("status"))
            break

    return templates.TemplateResponse(
        "job_detail.html",
        {
            "request": request,
            "job": job,
            "items": items,
            "auth_profiles": list_auth_profiles(),
            "last_status": last_status,
            "item_count": len(items),
            "next_runs": next_runs,
        },
    )


@frontend_router.get("/jobs/{job_id}/edit", response_class=HTMLResponse, dependencies=[Depends(require_login)])
def edit_job(request: Request, job_id: int) -> HTMLResponse:
    job = require_value(get_job(job_id), "任务不存在")
    return templates.TemplateResponse(
        "job_form.html",
        {
            "request": request,
            "job": job,
            "mail_profiles": list_mail_profiles(),
            "action": url_for_frontend(f"/jobs/{job_id}/edit"),
            "error": request.query_params.get("error", "")
        },
    )


@frontend_router.post("/jobs/{job_id}/edit", dependencies=[Depends(require_login)])
def update_job(job_id: int, form: dict[str, Any] = Depends(job_form)) -> RedirectResponse:
    from app.utils import validate_cron_expression
    from urllib.parse import quote
    import sqlite3
    old = require_value(get_job(job_id), "任务不存在")
    try:
        validate_cron_expression(form["cron_expression"])
    except ValueError as e:
        return frontend_redirect(f"/jobs/{job_id}/edit?error={quote(str(e))}")

    if not form.get("dingtalk_webhook"):
        form["dingtalk_webhook"] = old.get("dingtalk_webhook") or ""
    if not form.get("dingtalk_secret"):
        form["dingtalk_secret"] = old.get("dingtalk_secret") or ""
    try:
        save_job(form, job_id)
    except sqlite3.IntegrityError as e:
        if "UNIQUE constraint failed" in str(e) and "report_jobs.name" in str(e):
            return frontend_redirect(f"/jobs/{job_id}/edit?error={quote('任务名称已存在')}")
        return frontend_redirect(f"/jobs/{job_id}/edit?error={quote(f'数据库错误: {e}')}")
    reload_jobs()
    return frontend_redirect(f"/jobs/{job_id}")


@frontend_router.post("/jobs/{job_id}/delete", dependencies=[Depends(require_login)])
def remove_job(job_id: int) -> RedirectResponse:
    delete_job(job_id)
    reload_jobs()
    return frontend_redirect("/")


@frontend_router.post("/jobs/{job_id}/run", dependencies=[Depends(require_login)])
def run_job_now(job_id: int, background_tasks: BackgroundTasks) -> RedirectResponse:
    job = require_value(get_job(job_id), "任务不存在")
    run_id = create_run(job_id, str(job["name"]))
    background_tasks.add_task(run_job, job_id, run_id)
    return frontend_redirect(f"/runs/{run_id}?watch=1")


@frontend_router.get("/jobs/{job_id}/items/new", response_class=HTMLResponse, dependencies=[Depends(require_login)])
def new_item(request: Request, job_id: int) -> HTMLResponse:
    require_value(get_job(job_id), "任务不存在")
    return templates.TemplateResponse(
        "item_form.html",
        {
            "request": request,
            "item": {
                "job_id": job_id,
                "enabled": 1,
                "capture_mode": "viewport",
                "item_type": "grafana",
                "watermark_enabled": 0,
                "watermark_text": "{time}\n省客户服务中心\ndwangchengyi7(王诚毅)",
                "watermark_opacity": 65,
                "watermark_font_size": 24,
                "watermark_gap_x": 140,
                "watermark_gap_y": 140,
                "watermark_angle": -45,
            },
            "job_id": job_id,
            "auth_profiles": list_auth_profiles(),
            "template_sections": list_template_sections(),
            "action": url_for_frontend(f"/jobs/{job_id}/items/new"),
        },
    )


@frontend_router.post("/jobs/{job_id}/items/new", dependencies=[Depends(require_login)])
def create_item(job_id: int, form: dict[str, Any] = Depends(item_form)) -> RedirectResponse:
    form["job_id"] = job_id
    item_id = save_screenshot_item(form)
    return frontend_redirect(f"/jobs/{job_id}")


@frontend_router.get("/items/{item_id}/edit", response_class=HTMLResponse, dependencies=[Depends(require_login)])
def edit_item(request: Request, item_id: int) -> HTMLResponse:
    item = require_value(get_screenshot_item(item_id), "截图项不存在")
    return templates.TemplateResponse(
        "item_form.html",
        {
            "request": request,
            "item": item,
            "job_id": item["job_id"],
            "auth_profiles": list_auth_profiles(),
            "template_sections": list_template_sections(),
            "action": url_for_frontend(f"/items/{item_id}/edit"),
        },
    )


@frontend_router.post("/items/{item_id}/edit", dependencies=[Depends(require_login)])
def update_item(item_id: int, form: dict[str, Any] = Depends(item_form)) -> RedirectResponse:
    current = require_value(get_screenshot_item(item_id), "截图项不存在")
    form["job_id"] = current["job_id"]
    save_screenshot_item(form, item_id)
    return frontend_redirect(f"/jobs/{current['job_id']}")


@frontend_router.post("/items/{item_id}/delete", dependencies=[Depends(require_login)])
def remove_item(item_id: int) -> RedirectResponse:
    item = require_value(get_screenshot_item(item_id), "截图项不存在")
    delete_screenshot_item(item_id)
    return frontend_redirect(f"/jobs/{item['job_id']}")


@frontend_router.post("/items/{item_id}/test", dependencies=[Depends(require_login)])
def test_item(item_id: int, background_tasks: BackgroundTasks) -> RedirectResponse:
    require_value(get_screenshot_item(item_id), "截图项不存在")
    background_tasks.add_task(test_capture_item, item_id)
    return frontend_redirect("/runs?watch=1")


@frontend_router.get("/auth-profiles", dependencies=[Depends(require_login)])
def auth_profiles(request: Request) -> RedirectResponse:
    params = request.query_params
    query_str = f"?{params}" if params else ""
    return frontend_redirect(f"/#/auth{query_str}")


@frontend_router.get("/auth-profiles/new", response_class=HTMLResponse, dependencies=[Depends(require_login)])
def new_auth_profile(request: Request) -> HTMLResponse:
    return templates.TemplateResponse(
        "auth_form.html",
        {"request": request, "profile": {}, "action": url_for_frontend("/auth-profiles/new")},
    )


@frontend_router.post("/auth-profiles/new", dependencies=[Depends(require_login)])
def create_auth_profile(form: dict[str, Any] = Depends(auth_form)) -> RedirectResponse:
    auth_id = save_auth_profile(form)
    return frontend_redirect("/auth-profiles")


@frontend_router.get("/auth-profiles/{auth_id}/edit", response_class=HTMLResponse, dependencies=[Depends(require_login)])
def edit_auth_profile(request: Request, auth_id: int) -> HTMLResponse:
    profile = require_value(get_auth_profile(auth_id), "认证配置不存在")
    return templates.TemplateResponse(
        "auth_form.html",
        {"request": request, "profile": profile, "action": url_for_frontend(f"/auth-profiles/{auth_id}/edit")},
    )


@frontend_router.post("/auth-profiles/{auth_id}/edit", dependencies=[Depends(require_login)])
def update_auth_profile(auth_id: int, form: dict[str, Any] = Depends(auth_form)) -> RedirectResponse:
    old = get_auth_profile(auth_id)
    if old:
        # 若留空代表不修改，沿用旧凭据
        if not form.get("password"):
            form["password_secret"] = old.get("password_secret")
            
        if not form.get("username_value"):
            form["username_value"] = old.get("username_value")
            form["username"] = old.get("username")
            
        # 沿用未提交的高级选择器；如果页面提交了新值，则保存新配置。
        for field in ["username_selector", "password_selector", "submit_selector", "success_selector"]:
            if not form.get(field):
                form[field] = old.get(field) or form.get(field)

        changed = any(
            str(form.get(field) or "") != str(old.get(field) or "")
            for field in (
                "auth_type", "login_url", "username_value", "username", 
                "username_selector", "password_selector", "submit_selector", 
                "success_selector", "username_source", "password_source", 
                "username_env", "password_env", "storage_state_path"
            )
        )
        if form.get("password"):
            changed = True

        if changed:
            clear_auth_state_cache(dict(old))
            
    save_auth_profile(form, auth_id)
    return frontend_redirect("/auth-profiles")


@frontend_router.post("/auth-profiles/{auth_id}/delete", dependencies=[Depends(require_login)])
def remove_auth_profile(auth_id: int) -> RedirectResponse:
    old = get_auth_profile(auth_id)
    if old:
        clear_auth_state_cache(dict(old))
    delete_auth_profile(auth_id)
    return frontend_redirect("/auth-profiles")


@frontend_router.post("/auth-profiles/{auth_id}/test", dependencies=[Depends(require_login)])
def test_auth_profile_route(auth_id: int) -> RedirectResponse:
    from app.screenshot import test_auth_profile_login
    from urllib.parse import quote
    res = test_auth_profile_login(auth_id)
    if res["success"]:
        return frontend_redirect(f"/auth-profiles?success={quote(res['message'])}")
    return frontend_redirect(f"/auth-profiles?error={quote(res['message'])}")


@frontend_router.post("/auth-profiles/{auth_id}/test-async", dependencies=[Depends(require_login)])
def test_auth_profile_route_async(auth_id: int) -> dict[str, Any]:
    from app.screenshot import test_auth_profile_login
    return test_auth_profile_login(auth_id)


@frontend_router.get("/mail-profiles", dependencies=[Depends(require_login)])
def mail_profiles(request: Request) -> RedirectResponse:
    params = request.query_params
    query_str = f"?{params}" if params else ""
    return frontend_redirect(f"/#/mail{query_str}")


@frontend_router.get("/mail-profiles/new", response_class=HTMLResponse, dependencies=[Depends(require_login)])
def new_mail_profile(request: Request) -> HTMLResponse:
    return templates.TemplateResponse(
        "mail_form.html",
        {"request": request, "profile": {}, "action": url_for_frontend("/mail-profiles/new")},
    )


@frontend_router.post("/mail-profiles/new", dependencies=[Depends(require_login)])
def create_mail_profile(form: dict[str, Any] = Depends(mail_form)) -> RedirectResponse:
    mail_id = save_mail_profile(form)
    return frontend_redirect("/mail-profiles")


@frontend_router.get("/mail-profiles/{mail_id}/edit", response_class=HTMLResponse, dependencies=[Depends(require_login)])
def edit_mail_profile(request: Request, mail_id: int) -> HTMLResponse:
    profile = require_value(get_mail_profile(mail_id), "邮件配置不存在")
    return templates.TemplateResponse(
        "mail_form.html",
        {"request": request, "profile": profile, "action": url_for_frontend(f"/mail-profiles/{mail_id}/edit")},
    )


@frontend_router.post("/mail-profiles/{mail_id}/edit", dependencies=[Depends(require_login)])
def update_mail_profile(mail_id: int, form: dict[str, Any] = Depends(mail_form)) -> RedirectResponse:
    old = get_mail_profile(mail_id)
    if old:
        if not form.get("password"):
            form["password_secret"] = old.get("password_secret")
    save_mail_profile(form, mail_id)
    return frontend_redirect("/mail-profiles")


@frontend_router.post("/mail-profiles/{mail_id}/delete", dependencies=[Depends(require_login)])
def remove_mail_profile(mail_id: int) -> RedirectResponse:
    delete_mail_profile(mail_id)
    return frontend_redirect("/mail-profiles")


@frontend_router.post("/mail-profiles/{mail_id}/test", dependencies=[Depends(require_login)])
def test_mail_profile(mail_id: int, test_recipient: str = Form("")) -> RedirectResponse:
    profile = require_value(get_mail_profile(mail_id), "邮件配置不存在")
    recipient = test_recipient.strip()
    if not recipient:
        recipient = str(profile.get("recipients") or "").strip()
    if not recipient:
        recipient = str(profile.get("sender") or profile.get("username") or "").strip()
    if not recipient:
        from urllib.parse import quote
        return frontend_redirect(f"/mail-profiles?error={quote('收件人为空且无法提取默认收件人，请指定测试收件邮箱')}")
        
    import tempfile
    from docx import Document
    from urllib.parse import quote
    from app.mailer import send_report_mail

    temp_dir = tempfile.gettempdir()
    test_doc_path = os.path.join(temp_dir, "test_report.docx")
    try:
        doc = Document()
        doc.add_heading("邮件发送配置测试报告", level=1)
        doc.add_paragraph("这是一封由系统自动生成的巡检报告测试附件。")
        doc.add_paragraph("如果您收到这封邮件，说明您的 SMTP 邮件服务配置成功，且能够成功传递带附件的报告。")
        doc.save(test_doc_path)
        
        test_profile = dict(profile)
        test_profile["recipients"] = recipient
        test_profile["cc"] = ""
        
        dummy_job = {
            "name": "邮件配置连接测试",
            "environment": "SRE测试环境",
            "time_range_label": "即时"
        }
        send_report_mail(test_profile, dummy_job, test_doc_path)
    except Exception as exc:
        err_msg = f"{type(exc).__name__}: {exc}"
        app_logger.error(f"测试发送邮件失败: {err_msg}", exc_info=True)
        return frontend_redirect(f"/mail-profiles?error={quote(err_msg)}")
    finally:
        if os.path.exists(test_doc_path):
            try:
                os.remove(test_doc_path)
            except Exception:
                pass
                
    from urllib.parse import quote
    return frontend_redirect(f"/mail-profiles?success={quote('测试邮件已成功发出！请检查您的邮箱收件箱。')}")


@frontend_router.post("/mail-profiles/{mail_id}/test-async", dependencies=[Depends(require_login)])
def test_mail_profile_async(mail_id: int, test_recipient: str = Form("")) -> dict[str, Any]:
    profile = get_mail_profile(mail_id)
    if not profile:
        return {"success": False, "message": "邮件配置不存在"}
        
    recipient = test_recipient.strip()
    if not recipient:
        recipient = str(profile.get("recipients") or "").strip()
    if not recipient:
        recipient = str(profile.get("sender") or profile.get("username") or "").strip()
    if not recipient:
        return {"success": False, "message": "收件人为空且无法提取默认收件人，请指定测试收件邮箱"}
        
    import tempfile
    from docx import Document
    from app.mailer import send_report_mail

    temp_dir = tempfile.gettempdir()
    test_doc_path = os.path.join(temp_dir, "test_report.docx")
    try:
        doc = Document()
        doc.add_heading("邮件发送配置测试报告", level=1)
        doc.add_paragraph("这是一封由系统自动生成的巡检报告测试附件。")
        doc.add_paragraph("如果您收到这封邮件，说明您的 SMTP 邮件服务配置成功，且能够成功传递带附件的报告。")
        doc.save(test_doc_path)
        
        test_profile = dict(profile)
        test_profile["recipients"] = recipient
        test_profile["cc"] = ""
        
        dummy_job = {
            "name": "邮件配置连接测试",
            "environment": "SRE测试环境",
            "time_range_label": "即时"
        }
        send_report_mail(test_profile, dummy_job, test_doc_path)
        return {"success": True, "message": "测试邮件已成功发出！请检查您的邮箱收件箱。"}
    except Exception as exc:
        err_msg = f"{type(exc).__name__}: {exc}"
        app_logger.error(f"测试发送邮件失败: {err_msg}", exc_info=True)
        return {"success": False, "message": f"测试邮件发送失败: {err_msg}"}
    finally:
        if os.path.exists(test_doc_path):
            try:
                os.remove(test_doc_path)
            except Exception:
                pass


@frontend_router.get("/runs", dependencies=[Depends(require_login)])
def runs(request: Request) -> RedirectResponse:
    params = request.query_params
    query_str = f"?{params}" if params else ""
    return frontend_redirect(f"/#/runs{query_str}")


@frontend_router.get("/help", response_class=HTMLResponse, dependencies=[Depends(require_login)])
def help_page(request: Request) -> HTMLResponse:
    return templates.TemplateResponse(
        "help.html",
        {"request": request, "title": "配置说明与指南"},
    )



@frontend_router.get("/runs/{run_id}", response_class=HTMLResponse, dependencies=[Depends(require_login)])
def run_detail(request: Request, run_id: int) -> HTMLResponse:
    run = require_value(get_run(run_id), "运行记录不存在")
    results = [with_file_url(item) for item in list_screenshot_results(run_id)]
    if run.get("report_path"):
        filename = artifact_filename(run["report_path"])
        local_path = resolve_artifact_path(run["report_path"])
        if local_path.exists():
            run["report_url"] = get_artifact_url("reports", run_id, filename)
    return templates.TemplateResponse(
        "run_detail.html",
        {
            "request": request,
            "run": run,
            "results": results,
            "watch": request.query_params.get("watch", ""),
        },
    )


@frontend_router.get("/artifact/{kind}/{run_id}/{filename}", dependencies=[Depends(require_login)])
def artifact(kind: str, run_id: int, filename: str) -> FileResponse:
    base = {"screenshots": settings.screenshot_dir, "reports": settings.report_dir}.get(kind)
    if base is None:
        raise HTTPException(status_code=404, detail="未知文件类型")
        
    yyyy, mm, dd = get_run_date(run_id)
    new_path = (base / yyyy / mm / dd / f"run-{run_id}" / filename).resolve()
    legacy_path = (base / str(run_id) / filename).resolve()
    
    if new_path.exists() and is_under(new_path, base.resolve()):
        return FileResponse(new_path)
    elif legacy_path.exists() and is_under(legacy_path, base.resolve()):
        return FileResponse(legacy_path)
        
    raise HTTPException(status_code=404, detail="文件不存在")


def require_value(value: Any, message: str) -> Any:
    if value is None:
        raise HTTPException(status_code=404, detail=message)
    return value


def with_file_url(item: dict[str, Any]) -> dict[str, Any]:
    file_path = item.get("file_path") or ""
    run_id = item.get("run_id")
    if file_path and run_id:
        filename = artifact_filename(file_path)
        local_path = resolve_artifact_path(file_path)
        if local_path.exists():
            item["file_url"] = get_artifact_url("screenshots", run_id, filename)
    return item


def is_under(path: Path, base: Path) -> bool:
    try:
        path.relative_to(base)
        return True
    except ValueError:
        return False


def parse_optional_int(value: str | None) -> int | None:
    if value is None or value == "":
        return None
    try:
        return int(value)
    except (ValueError, TypeError) as e:
        raise HTTPException(status_code=400, detail="参数格式错误") from e


# ==============================================================================
# 周期报告与数据库重置路由
# ==============================================================================

@frontend_router.post("/settings/reset", dependencies=[Depends(require_login)])
def reset_database() -> RedirectResponse:
    from app.db import reset_database_data
    from app.scheduler import reload_jobs
    
    app_logger.info("用户请求一键重置所有本地测试数据")
    try:
        reset_database_data()
        reload_jobs()
    except Exception as exc:
        app_logger.error(f"重置数据库失败: {exc}", exc_info=True)
        
    from urllib.parse import quote
    return frontend_redirect(f"/?success={quote('系统所有配置及历史数据已成功清空并重置！')}")


@frontend_router.get("/periodic-reports", dependencies=[Depends(require_login)])
def periodic_reports_list(request: Request) -> RedirectResponse:
    params = request.query_params
    query_str = f"?{params}" if params else ""
    return frontend_redirect(f"/#/periodic{query_str}")


@frontend_router.get("/periodic-reports/{report_type}/edit", response_class=HTMLResponse, dependencies=[Depends(require_login)])
def edit_periodic_report(request: Request, report_type: str) -> HTMLResponse:
    from app.repository import get_periodic_report_setting, list_mail_profiles
    from app.utils import decrypt_secret
    setting = get_periodic_report_setting(report_type)
    if not setting:
        raise HTTPException(status_code=404, detail="周期报告配置不存在")
        
    # 解密敏感字段以在表单中回显
    for f in ("dingtalk_webhook", "dingtalk_secret"):
        if setting.get(f):
            try:
                setting[f] = decrypt_secret(setting[f])
            except Exception:
                pass
        
    return templates.TemplateResponse(
        "periodic_form.html",
        {
            "request": request,
            "setting": setting,
            "mail_profiles": list_mail_profiles(),
            "action": url_for_frontend(f"/periodic-reports/{report_type}/edit")
        }
    )


@frontend_router.post("/periodic-reports/{report_type}/edit", dependencies=[Depends(require_login)])
def update_periodic_report(
    report_type: str,
    enabled: str | None = Form(None),
    name: str = Form(...),
    schedule_mode: str = Form("simple"),
    cron_expression: str = Form(""),
    mail_profile_id: str | None = Form(None),
    include_docx: str | None = Form(None),
    include_screenshots: str | None = Form(None),
    send_empty_report: str | None = Form(None),
    wait_for_daily_jobs: str | None = Form(None),
    retry_until_time: str = Form("23:50"),
    retry_interval_minutes: int = Form(10),
    recipients_override: str = Form(""),
    frequency: str = Form("weekly"),
    day_of_week: str | None = Form(None),
    day_of_month: str | None = Form(None),
    ampm: list[str] = Form([]),
    hour: list[int] = Form([]),
    minute: list[int] = Form([]),
    send_on_timeout: int = Form(1),
    enable_email: str | None = Form(None),
    dingtalk_enabled: str | None = Form(None),
    dingtalk_webhook: str = Form(""),
    dingtalk_secret: str = Form(""),
    dingtalk_keyword: str = Form(""),
) -> RedirectResponse:
    import json
    from app.utils import times_to_cron, validate_cron_expression
    from app.repository import save_periodic_report_setting
    from app.scheduler import reload_jobs
    from urllib.parse import quote

    times_list = []
    for i in range(len(hour)):
        times_list.append({
            "ampm": ampm[i] if i < len(ampm) else "am",
            "hour": int(hour[i]),
            "minute": int(minute[i]) if i < len(minute) else 0
        })

    if schedule_mode == "simple":
        if frequency == "monthly_last":
            cron_list = []
            label_times = []
            for t in times_list:
                t_ampm = t.get("ampm", "am")
                t_hour = int(t.get("hour", 9))
                t_minute = int(t.get("minute", 0))
                
                if t_ampm == "pm" and t_hour < 12:
                    h_24 = t_hour + 12
                elif t_ampm == "am" and t_hour == 12:
                    h_24 = 0
                else:
                    h_24 = t_hour
                cron_list.append(f"0 {t_minute} {h_24} * * *")
                t_ampm_cn = "上午" if t_ampm == "am" else "下午"
                label_times.append(f"{t_ampm_cn}{t_hour}点{t_minute:02d}分" if t_minute else f"{t_ampm_cn}{t_hour}点")
            
            cron_expr = ";".join(cron_list)
            label_expr = f"每月最后一天 {', '.join(label_times)}"
            schedule_config = json.dumps({
                "frequency": "monthly_last",
                "times": times_list
            })
        else:
            cron_expr, label_expr = times_to_cron(frequency, day_of_week, day_of_month, times_list)
            schedule_config = json.dumps({
                "frequency": frequency,
                "day_of_week": day_of_week,
                "day_of_month": day_of_month,
                "times": times_list
            })
    else:
        cron_expr = cron_expression
        label_expr = "高级 Cron"
        schedule_config = ""

    try:
        validate_cron_expression(cron_expr)
    except ValueError as e:
        return frontend_redirect(f"/periodic-reports/{report_type}/edit?error={quote(str(e))}")

    update_data = {
        "enabled": 1 if enabled else 0,
        "name": name,
        "schedule_mode": schedule_mode,
        "schedule_label": label_expr,
        "cron_expression": cron_expr,
        "mail_profile_id": parse_optional_int(mail_profile_id),
        "include_screenshots": 1 if include_screenshots else 0,
        "include_docx": 1 if include_docx else 0,
        "send_empty_report": 1 if send_empty_report else 0,
        "wait_for_daily_jobs": 1 if wait_for_daily_jobs else 0,
        "retry_until_time": retry_until_time,
        "retry_interval_minutes": retry_interval_minutes,
        "recipients_override": recipients_override,
        "schedule_config": schedule_config,
        "send_on_timeout": send_on_timeout,
        "enable_email": 1 if enable_email else 0,
        "dingtalk_enabled": 1 if dingtalk_enabled else 0,
        "dingtalk_webhook": dingtalk_webhook.strip(),
        "dingtalk_secret": dingtalk_secret.strip(),
        "dingtalk_keyword": dingtalk_keyword.strip(),
    }
    
    save_periodic_report_setting(report_type, update_data)
    reload_jobs()
    return frontend_redirect(f"/periodic-reports?success={quote('周期汇总配置保存成功！')}")


@frontend_router.post("/periodic-reports/{report_type}/run", dependencies=[Depends(require_login)])
def run_periodic_report_now(report_type: str, background_tasks: BackgroundTasks) -> RedirectResponse:
    from app.periodic_reporter import trigger_periodic_report
    from urllib.parse import quote
    background_tasks.add_task(trigger_periodic_report, report_type, True)
    return frontend_redirect(f"/periodic-reports?success={quote('周期汇总报告生成任务已在后台启动，打包发信中，请稍候刷新页面查看。')}")


@frontend_router.get("/periodic-reports/download/{run_id}", dependencies=[Depends(require_login)])
def download_periodic_report_zip(run_id: int) -> FileResponse:
    from app.repository import get_periodic_report_run
    run = get_periodic_report_run(run_id)
    if not run or not run.get("zip_path"):
        raise HTTPException(status_code=404, detail="周期报告包不存在")
    path = Path(run["zip_path"]).resolve()
    if not path.exists():
        raise HTTPException(status_code=400, detail="周期报告附件已按数据保留策略被自动清理")
    return FileResponse(path, media_type="application/zip", filename=path.name)


@frontend_router.post("/periodic-reports/delete/{run_id}", dependencies=[Depends(require_login)])
def delete_periodic_report_run_route(run_id: int) -> RedirectResponse:
    from app.repository import get_periodic_report_run, delete_periodic_report_run
    run = get_periodic_report_run(run_id)
    if run:
        z_path = run.get("zip_path")
        if z_path and os.path.exists(z_path):
            try:
                os.remove(z_path)
            except Exception:
                pass
        delete_periodic_report_run(run_id)
    return frontend_redirect("/periodic-reports")


# ==============================================================================
# 存储管理与数据保留策略 API 路由
# ==============================================================================

@api_router.get("/storage", dependencies=[Depends(require_api_login)])
def api_get_storage() -> dict[str, Any]:
    from app.cleanup import get_storage_usage
    from app.db import connect
    
    try:
        usage = get_storage_usage()
    except Exception as e:
        app_logger.error(f"获取存储占用统计失败: {e}", exc_info=True)
        usage = {
            "total_bytes": 0, "screenshots_bytes": 0, "reports_bytes": 0,
            "logs_bytes": 0, "browser_state_bytes": 0, "periodic_reports_bytes": 0,
            "sqlite_db_bytes": 0, "estimated_cleanup_bytes": 0, "estimated_cleanup_files": 0
        }
    
    with connect() as conn:
        row = conn.execute("SELECT * FROM storage_cleanup_settings ORDER BY id DESC LIMIT 1").fetchone()
        config = dict(row) if row else {}
    
    with connect() as conn:
        row_run = conn.execute("SELECT * FROM storage_cleanup_runs ORDER BY id DESC LIMIT 1").fetchone()
        latest_run = dict(row_run) if row_run else None
        
    return {
        "usage": usage,
        "config": config,
        "latest_run": latest_run
    }


@api_router.put("/storage/settings", dependencies=[Depends(require_api_login)])
def api_update_storage_settings(payload: dict[str, Any]) -> dict[str, Any]:
    from app.db import connect
    from app.scheduler import reload_jobs
    import json
    
    enabled = bool_value(payload.get("enabled"), True)
    allow_manual_cleanup = bool_value(payload.get("allow_manual_cleanup"), True)
    cleanup_schedule_mode = text_value(payload, "cleanup_schedule_mode", "daily")
    
    cleanup_schedule_label = text_value(payload, "cleanup_schedule_label", "")
    cleanup_schedule_config = text_value(payload, "cleanup_schedule_config", "")
    
    if not cleanup_schedule_label:
        if cleanup_schedule_config:
            try:
                cfg = json.loads(cleanup_schedule_config)
                time_str = cfg.get("time") or "02:30"
                if cleanup_schedule_mode == "daily":
                    cleanup_schedule_label = f"每天 {time_str}"
                elif cleanup_schedule_mode == "weekly":
                    dow_val = str(cfg.get("day_of_week") or "1")
                    mapping = {"1": "周一", "2": "周二", "3": "周三", "4": "周四", "5": "周五", "6": "周六", "7": "周日", "0": "周日"}
                    cleanup_schedule_label = f"每周 {mapping.get(dow_val, '周一')} {time_str}"
                elif cleanup_schedule_mode == "monthly":
                    day_val = str(cfg.get("day_of_month") or "1")
                    cleanup_schedule_label = f"每月 {day_val}号 {time_str}"
            except Exception:
                cleanup_schedule_label = "每天 02:30"
        else:
            cleanup_schedule_label = "每天 02:30"
            
    periodic_sent_retention_days = int_value(payload, "periodic_sent_retention_days", 3)
    periodic_failed_retention_days = int_value(payload, "periodic_failed_retention_days", 30)
    screenshot_retention_days = int_value(payload, "screenshot_retention_days", 7)
    report_retention_days = int_value(payload, "report_retention_days", 30)
    run_record_retention_days = int_value(payload, "run_record_retention_days", 90)
    log_retention_days = int_value(payload, "log_retention_days", 14)
    browser_state_retention_days = int_value(payload, "browser_state_retention_days", 30)
    browser_state_cleanup_enabled = bool_value(payload.get("browser_state_cleanup_enabled"), False)
    protect_recent_days = int_value(payload, "protect_recent_days", 3)
    
    with connect() as conn:
        row = conn.execute("SELECT id FROM storage_cleanup_settings ORDER BY id DESC LIMIT 1").fetchone()
        if row:
            setting_id = row[0]
            conn.execute(
                """
                UPDATE storage_cleanup_settings SET
                    enabled = ?, allow_manual_cleanup = ?, cleanup_schedule_mode = ?,
                    cleanup_schedule_label = ?, cleanup_schedule_config = ?,
                    periodic_sent_retention_days = ?, periodic_failed_retention_days = ?,
                    screenshot_retention_days = ?, report_retention_days = ?,
                    run_record_retention_days = ?, log_retention_days = ?,
                    browser_state_retention_days = ?, browser_state_cleanup_enabled = ?,
                    protect_recent_days = ?, updated_at = CURRENT_TIMESTAMP
                WHERE id = ?
                """,
                (
                    enabled, allow_manual_cleanup, cleanup_schedule_mode,
                    cleanup_schedule_label, cleanup_schedule_config,
                    periodic_sent_retention_days, periodic_failed_retention_days,
                    screenshot_retention_days, report_retention_days,
                    run_record_retention_days, log_retention_days,
                    browser_state_retention_days, browser_state_cleanup_enabled,
                    protect_recent_days, setting_id
                )
            )
        else:
            conn.execute(
                """
                INSERT INTO storage_cleanup_settings (
                    enabled, allow_manual_cleanup, cleanup_schedule_mode,
                    cleanup_schedule_label, cleanup_schedule_config,
                    periodic_sent_retention_days, periodic_failed_retention_days,
                    screenshot_retention_days, report_retention_days,
                    run_record_retention_days, log_retention_days,
                    browser_state_retention_days, browser_state_cleanup_enabled,
                    protect_recent_days
                ) VALUES (?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?)
                """,
                (
                    enabled, allow_manual_cleanup, cleanup_schedule_mode,
                    cleanup_schedule_label, cleanup_schedule_config,
                    periodic_sent_retention_days, periodic_failed_retention_days,
                    screenshot_retention_days, report_retention_days,
                    run_record_retention_days, log_retention_days,
                    browser_state_retention_days, browser_state_cleanup_enabled,
                    protect_recent_days
                )
            )
    
    reload_jobs()
    return {"ok": True}


@api_router.post("/storage/estimate", dependencies=[Depends(require_api_login)])
def api_estimate_storage() -> dict[str, Any]:
    from app.cleanup import estimate_cleanup
    try:
        res = estimate_cleanup()
        return {"ok": True, "result": res}
    except Exception as e:
        app_logger.error(f"执行预估清理失败: {e}", exc_info=True)
        raise HTTPException(status_code=500, detail=f"预估清理失败: {e}")


@api_router.post("/storage/cleanup", dependencies=[Depends(require_api_login)])
def api_run_storage_cleanup() -> dict[str, Any]:
    from app.cleanup import run_cleanup
    try:
        res = run_cleanup(mode="manual", dry_run=False)
        if res.get("status") == "failed" and "系统策略限制" in res.get("error_summary", ""):
            raise HTTPException(status_code=400, detail=res["error_summary"])
        return {"ok": True, "result": res}
    except HTTPException as he:
        raise he
    except Exception as e:
        app_logger.error(f"执行物理清理失败: {e}", exc_info=True)
        raise HTTPException(status_code=500, detail=f"执行清理失败: {e}")


@api_router.get("/storage/cleanup-runs", dependencies=[Depends(require_api_login)])
def api_get_cleanup_runs() -> dict[str, Any]:
    from app.db import connect
    with connect() as conn:
        rows = conn.execute("SELECT * FROM storage_cleanup_runs ORDER BY id DESC LIMIT 50").fetchall()
        runs = [dict(row) for row in rows]
    return {"runs": runs}


app.include_router(frontend_router)
app.include_router(api_router)
