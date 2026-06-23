import docx
from pathlib import Path

doc_path = Path("服务器巡检-yyyy-mm-dd.docx")
if doc_path.exists():
    doc = docx.Document(doc_path)
    print("=== 段落列表 ===")
    for i, p in enumerate(doc.paragraphs):
        if p.text.strip():
            print(f"[{i}]: {p.text}")
            
    print("\n=== 表格列表 ===")
    for t_idx, table in enumerate(doc.tables):
        print(f"\n表格 {t_idx}:")
        for r_idx, row in enumerate(table.rows):
            cells = [c.text.strip().replace('\n', ' ') for c in row.cells]
            print(f"  行 {r_idx}: {cells}")
else:
    print("模板文件不存在")
